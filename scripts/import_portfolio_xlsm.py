"""
Скрипт импорта реального портфеля проектов автоматизации из Портфель_проектов_1.0.14h.xlsm в PostgreSQL.
"""

import os
import sys
from pathlib import Path
from datetime import date, datetime
import asyncio
import openpyxl

# Добавляем корень проекта и backend в sys.path
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))
BACKEND_DIR = ROOT_DIR / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from sqlalchemy import select, delete
from app.core.database import AsyncSessionLocal
from app.models.project import PortfolioProject, ProjectStatus
from app.models.task import PortfolioTask, TaskStatus, TaskType
from app.services.cpm.engine import CPMEngine
from app.services.cpm.types import CPMNode, CPMEdge
from app.services.cpm.calendar import WorkCalendar


FILE_NAME = "Портфель_проектов_1.0.14h.xlsm"


def coerce_date(val, default_date=None) -> date:
    """Приведение ячейки Excel к date."""
    if val is None or str(val).strip() in ("", "None", "nan"):
        return default_date or date.today()
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
            try:
                return datetime.strptime(val.strip(), fmt).date()
            except ValueError:
                pass
    return default_date or date.today()


def coerce_int(val, default=1) -> int:
    try:
        return max(0, int(float(val)))
    except (ValueError, TypeError):
        return default


async def import_portfolio():
    excel_path = ROOT_DIR / FILE_NAME
    if not excel_path.exists():
        print(f"❌ Файл {FILE_NAME} не найден в {ROOT_DIR}")
        return

    print(f"📂 Загружаем книгу: {excel_path}...")
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)

    # 1. Извлекаем список проектов из листа 'Проекты'
    ws_projects = wb['Проекты']
    raw_projects = []
    for row in ws_projects.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        # Колонка 5 (индекс 4) - полное наименование, колонка 6 - руководитель
        name = str(row[4]).strip() if len(row) > 4 and row[4] else f"Проект {code}"
        manager = str(row[5]).strip() if len(row) > 5 and row[5] else "Не назначен"
        priority = str(row[7]).strip() if len(row) > 7 and row[7] else "Средний"
        raw_projects.append({
            "code": code,
            "name": name,
            "manager": manager,
            "priority": priority,
        })

    print(f"✅ Найдено проектов в реестре: {len(raw_projects)}")

    # 2. Извлекаем задачи/этапы из листа '03_M06_Данные'
    ws_tasks = wb['03_M06_Данные']
    # Заголовки: Проект (0), ID (1), Порядок (2), Ответственный (3), Этап/задача (4), План старт (5), Длительность (6), План финиш (7)
    tasks_by_project = {}
    for row in ws_tasks.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        proj_code = str(row[0]).strip()
        stage_id = str(row[1]).strip() if row[1] is not None else "1"
        sort_order = coerce_int(row[2], default=0) if len(row) > 2 else 0
        assignee = str(row[3]).strip() if len(row) > 3 and row[3] else None
        task_name = str(row[4]).strip() if len(row) > 4 and row[4] else f"Этап {stage_id}"
        p_start = coerce_date(row[5]) if len(row) > 5 else date.today()
        duration = coerce_int(row[6], default=1) if len(row) > 6 else 1
        p_finish = coerce_date(row[7], default_date=p_start) if len(row) > 7 else p_start

        if proj_code not in tasks_by_project:
            tasks_by_project[proj_code] = []

        tasks_by_project[proj_code].append({
            "stage_id": stage_id,
            "sort_order": sort_order,
            "assignee": assignee,
            "name": task_name,
            "planned_start": p_start,
            "planned_duration": duration,
            "planned_finish": p_finish,
        })

    total_tasks_count = sum(len(t) for t in tasks_by_project.values())
    print(f"✅ Найдено этапов в M06: {total_tasks_count}")

    # 3. Сохранение в PostgreSQL через сессию
    cpm_engine = CPMEngine(calendar=WorkCalendar())

    async with AsyncSessionLocal() as session:
        print("\n💾 Сохраняем данные в PostgreSQL...")
        
        for p_info in raw_projects:
            code = p_info["code"]
            proj_tasks = tasks_by_project.get(code, [])

            # Вычисляем директивные даты проекта на основе его задач
            min_date = min((t["planned_start"] for t in proj_tasks), default=date.today())
            max_date = max((t["planned_finish"] for t in proj_tasks), default=date.today())

            # Поиск или создание проекта
            stmt = select(PortfolioProject).where(PortfolioProject.code == code)
            res = await session.execute(stmt)
            project = res.scalar_one_or_none()

            raci_matrix = {
                "pm": p_info["manager"],
                "sponsor": "Дирекция по автоматизации",
                "priority": p_info["priority"],
            }
            passport_data = {
                "priority": p_info["priority"],
                "industry": "Мебельное и деревообрабатывающее производство",
            }

            if not project:
                project = PortfolioProject(
                    code=code,
                    name=p_info["name"],
                    description=f"Проект автоматизации производства: {p_info['name']}",
                    status=ProjectStatus.EXECUTION,
                    planned_start_date=min_date,
                    planned_finish_date=max_date,
                    raci_matrix=raci_matrix,
                    passport_data=passport_data,
                )
                session.add(project)
                await session.flush()
            else:
                project.name = p_info["name"]
                project.planned_start_date = min_date
                project.planned_finish_date = max_date
                project.raci_matrix = raci_matrix
                # Удаляем старые задачи проекта
                await session.execute(delete(PortfolioTask).where(PortfolioTask.project_id == project.id))
                await session.flush()

            # Добавляем задачи проекта с иерархией WBS
            db_tasks = []
            for idx, t_info in enumerate(proj_tasks, start=1):
                wbs_code = f"1.{idx}"
                task_type = TaskType.MILESTONE if t_info["planned_duration"] == 0 else TaskType.TASK
                
                task = PortfolioTask(
                    project_id=project.id,
                    wbs_code=wbs_code,
                    parent_wbs="1.0",
                    level=2,
                    sort_order=t_info["sort_order"] or idx,
                    name=t_info["name"],
                    task_type=task_type,
                    status=TaskStatus.IN_PROGRESS if idx <= 2 else TaskStatus.NOT_STARTED,
                    planned_duration=t_info["planned_duration"],
                    planned_start=t_info["planned_start"],
                    planned_finish=t_info["planned_finish"],
                    actual_start=t_info["planned_start"] if idx <= 2 else None,
                    progress_percent=100.0 if idx == 1 else (50.0 if idx == 2 else 0.0),
                    is_plan_locked=(idx <= 2),
                    raci_assignee=t_info["assignee"] or p_info["manager"],
                    raci_approver=p_info["manager"],
                )
                session.add(task)
                db_tasks.append(task)

            await session.flush()

            # Запускаем CPM Engine для каждого проекта
            cpm_nodes = [
                CPMNode(
                    id=t.id,
                    wbs_code=t.wbs_code,
                    name=t.name,
                    duration_days=t.planned_duration,
                    task_type=t.task_type,
                    planned_start=t.planned_start,
                    planned_finish=t.planned_finish,
                    actual_start=t.actual_start,
                    progress_percent=t.progress_percent,
                )
                for t in db_tasks
            ]

            cpm_result = cpm_engine.calculate(
                project_start=project.planned_start_date,
                nodes=cpm_nodes,
                edges=[],
            )

            # Обновляем расчетные даты
            nodes_map = {n.id: n for n in cpm_result.tasks}
            for t in db_tasks:
                calc = nodes_map.get(t.id)
                if calc:
                    t.early_start = calc.early_start
                    t.early_finish = calc.early_finish
                    t.late_start = calc.late_start
                    t.late_finish = calc.late_finish
                    t.total_float = calc.total_float
                    t.free_float = calc.free_float
                    t.is_critical = calc.is_critical

            print(f"   📁 Проект [{code}] '{p_info['name'][:40]}...' -> Загружено этапов: {len(db_tasks)}")

        await session.commit()

    print("\n🎉 ВСЕ ПРОЕКТЫ И ЭТАПЫ УСПЕШНО ИМПОРТИРОВАНЫ В POSTGRESQL!")


if __name__ == "__main__":
    asyncio.run(import_portfolio())