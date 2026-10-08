import uuid
import os
import shutil
import tempfile
from datetime import date, datetime
from typing import List, Optional, Any
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, status
from pydantic import BaseModel
from sqlalchemy import select, desc
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl

from app.api.deps import get_db, get_current_user, require_permission
from app.models.auth import User
from app.models.diary import DiaryEntry, DiaryEntryType
from app.models.task import PortfolioTask
from app.models.project import PortfolioProject

router = APIRouter()


def coerce_diary_date(val: Any, default_date: Optional[date] = None) -> date:
    if val is None or str(val).strip() in ("", "None", "nan"):
        return default_date or date.today()
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(val.strip(), fmt).date()
            except ValueError:
                pass
    return default_date or date.today()

class DiaryEntryCreateSchema(BaseModel):
    category: Optional[str] = "B"
    wbs_affected: Optional[str] = None
    task_id: Optional[uuid.UUID] = None
    entry_date: Optional[date] = None
    line_section: Optional[str] = None
    content: str
    solution: Optional[str] = None
    physical_action: Optional[str] = None
    purpose: Optional[str] = None
    author_name: Optional[str] = None
    deadline_str: Optional[str] = None
    status: Optional[str] = "В работе"
    source: Optional[str] = "Оперативка"

class DiaryEntryUpdateSchema(BaseModel):
    category: Optional[str] = None
    wbs_affected: Optional[str] = None
    task_id: Optional[uuid.UUID] = None
    entry_date: Optional[date] = None
    line_section: Optional[str] = None
    content: Optional[str] = None
    solution: Optional[str] = None
    physical_action: Optional[str] = None
    purpose: Optional[str] = None
    author_name: Optional[str] = None
    deadline_str: Optional[str] = None
    status: Optional[str] = None
    source: Optional[str] = None

@router.get("/diary")
@router.get("/projects/{project_id}/diary")
async def list_diary_entries(
    project_id: Optional[uuid.UUID] = None,
    category: Optional[str] = None,
    line_section: Optional[str] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission("projects:view")),
):
    stmt = select(DiaryEntry).order_by(desc(DiaryEntry.entry_date), desc(DiaryEntry.created_at))
    if project_id:
        stmt = stmt.where(DiaryEntry.project_id == project_id)
    if category:
        stmt = stmt.where(DiaryEntry.category == category.upper())
    if line_section:
        stmt = stmt.where(DiaryEntry.line_section == line_section)
    if status:
        stmt = stmt.where(DiaryEntry.status == status)
    if search:
        s = f"%{search}%"
        stmt = stmt.where(
            (DiaryEntry.content.ilike(s)) |
            (DiaryEntry.solution.ilike(s)) |
            (DiaryEntry.author_name.ilike(s)) |
            (DiaryEntry.line_section.ilike(s)) |
            (DiaryEntry.wbs_affected.ilike(s))
        )
    res = await db.execute(stmt)
    return res.scalars().all()

@router.post("/projects/{project_id}/diary")
async def create_project_diary_entry(
    project_id: uuid.UUID,
    payload: DiaryEntryCreateSchema,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    entry = DiaryEntry(
        project_id=project_id,
        task_id=payload.task_id,
        category=payload.category or "B",
        wbs_affected=payload.wbs_affected,
        entry_date=payload.entry_date or date.today(),
        line_section=payload.line_section,
        content=payload.content,
        solution=payload.solution,
        physical_action=payload.physical_action,
        purpose=payload.purpose,
        author_name=payload.author_name or current_user.full_name,
        deadline_str=payload.deadline_str,
        status=payload.status or "В работе",
        source=payload.source or "Оперативка",
        entry_type=DiaryEntryType.DAILY_LOG,
    )
    db.add(entry)
    await db.commit()
    await db.refresh(entry)
    return entry

@router.get("/tasks/{task_id}/diary")
async def get_task_diary_entries(
    task_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission("projects:view")),
):
    stmt = select(DiaryEntry).where(DiaryEntry.task_id == task_id).order_by(desc(DiaryEntry.entry_date))
    res = await db.execute(stmt)
    return res.scalars().all()

@router.post("/tasks/{task_id}/diary")
async def create_task_diary_entry(
    task_id: uuid.UUID,
    payload: DiaryEntryCreateSchema,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    task = await db.get(PortfolioTask, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="Задача не найдена")

    entry = DiaryEntry(
        project_id=task.project_id,
        task_id=task_id,
        category=payload.category or "B",
        wbs_affected=payload.wbs_affected or task.wbs_code,
        entry_date=payload.entry_date or date.today(),
        line_section=payload.line_section,
        content=payload.content,
        solution=payload.solution,
        physical_action=payload.physical_action,
        purpose=payload.purpose,
        author_name=payload.author_name or current_user.full_name,
        deadline_str=payload.deadline_str,
        status=payload.status or "В работе",
        source=payload.source or "Оперативка",
        entry_type=DiaryEntryType.DAILY_LOG,
    )
    db.add(entry)
    await db.commit()
    await db.refresh(entry)
    return entry

@router.patch("/diary/{entry_id}")
@router.patch("/tasks/{task_id}/diary/{entry_id}")
async def update_diary_entry(
    entry_id: uuid.UUID,
    payload: DiaryEntryUpdateSchema,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    entry = await db.get(DiaryEntry, entry_id)
    if not entry:
        raise HTTPException(status_code=404, detail="Запись дневника не найдена")
    data = payload.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(entry, k, v)
    await db.commit()
    await db.refresh(entry)
    return entry

@router.delete("/diary/{entry_id}")
@router.delete("/tasks/{task_id}/diary/{entry_id}")
async def delete_diary_entry(
    entry_id: uuid.UUID,
    task_id: Optional[uuid.UUID] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission("diary:edit")),
):
    entry = await db.get(DiaryEntry, entry_id)
    if not entry or (task_id and entry.task_id != task_id):
        raise HTTPException(status_code=404, detail="Запись дневника не найдена")
    await db.delete(entry)
    await db.commit()
    return {"status": "deleted"}


@router.post("/diary/import-excel")
async def import_diary_excel(
    file: UploadFile = File(...),
    project_id: Optional[uuid.UUID] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission("projects:view")),
):
    """
    Импорт записей цехового дневника производства работ из книги Excel.
    """
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Поддерживаются только форматы .xlsx и .xlsm",
        )

    suffix = ".xlsm" if file.filename.lower().endswith(".xlsm") else ".xlsx"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        shutil.copyfileobj(file.file, tmp)
        tmp_path = tmp.name

    imported_count = 0
    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
        # 1. Поиск целевого листа
        target_ws = None
        for name in wb.sheetnames:
            lower = name.lower()
            if any(k in lower for k in ("дневник", "журнал", "diary", "оперативк")):
                target_ws = wb[name]
                break
        if not target_ws:
            target_ws = wb.active

        # 2. Поиск строки заголовков
        header_row_idx = None
        col_map = {}
        for r_idx, row in enumerate(target_ws.iter_rows(values_only=True, max_row=15), start=1):
            row_lower = [str(c).strip().lower() if c is not None else "" for c in row]
            matches = sum(1 for c in row_lower if any(k in c for k in ("задач", "wbs", "участ", "решен", "проект", "категор")))
            if matches >= 2:
                header_row_idx = r_idx
                for c_idx, col_name in enumerate(row_lower):
                    if not col_name:
                        continue
                    if any(k in col_name for k in ("проект", "код проекта")):
                        col_map["project"] = c_idx
                    elif any(k in col_name for k in ("кат", "категор")):
                        col_map["category"] = c_idx
                    elif any(k in col_name for k in ("wbs", "влиян", "этап")):
                        col_map["wbs"] = c_idx
                    elif any(k in col_name for k in ("дат", "date")):
                        col_map["date"] = c_idx
                    elif any(k in col_name for k in ("участ", "линия", "оборудован")):
                        col_map["section"] = c_idx
                    elif any(k in col_name for k in ("задач", "проблем", "инцидент", "содержан", "работ")):
                        col_map["task"] = c_idx
                    elif any(k in col_name for k in ("решен", "solution")):
                        col_map["solution"] = c_idx
                    elif any(k in col_name for k in ("физическ", "действие", "что сделать")):
                        col_map["physical"] = c_idx
                    elif any(k in col_name for k in ("зачем", "цель", "обоснован")):
                        col_map["purpose"] = c_idx
                    elif any(k in col_name for k in ("ответств", "автор", "фио")):
                        col_map["author"] = c_idx
                    elif any(k in col_name for k in ("срок", "дедлайн")):
                        col_map["deadline"] = c_idx
                    elif any(k in col_name for k in ("статус", "status")):
                        col_map["status"] = c_idx
                    elif any(k in col_name for k in ("источник", "source")):
                        col_map["source"] = c_idx
                break

        if not header_row_idx or "task" not in col_map:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="В таблице Excel не найдена колонка 'Задача' или заголовок строк дневника",
            )

        # Подгружаем все проекты для сопоставления
        all_projs = (await db.execute(select(PortfolioProject))).scalars().all()
        proj_by_code = {p.code.upper().strip(): p for p in all_projs}
        fallback_proj = all_projs[0] if all_projs else None

        # 3. Построчное чтение и создание записей
        for row in target_ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not row or len(row) <= col_map["task"]:
                continue
            task_text = str(row[col_map["task"]]).strip() if row[col_map["task"]] is not None else ""
            if not task_text or task_text in ("—", "-", "None"):
                continue

            target_project = None
            if "project" in col_map and row[col_map["project"]] is not None:
                p_code = str(row[col_map["project"]]).upper().strip()
                target_project = proj_by_code.get(p_code)

            if not target_project and project_id:
                target_project = await db.get(PortfolioProject, project_id)

            if not target_project:
                target_project = fallback_proj

            if not target_project:
                continue

            wbs_val = str(row[col_map["wbs"]]).strip() if "wbs" in col_map and row[col_map["wbs"]] is not None else None
            matched_task = None
            if wbs_val:
                t_stmt = select(PortfolioTask).where(
                    PortfolioTask.project_id == target_project.id,
                    PortfolioTask.wbs_code.ilike(f"%{wbs_val}%")
                )
                matched_task = (await db.execute(t_stmt)).scalars().first()

            entry = DiaryEntry(
                project_id=target_project.id,
                task_id=matched_task.id if matched_task else None,
                category=str(row[col_map.get("category", -1)] or "B").upper().strip()[:10] if "category" in col_map and row[col_map["category"]] is not None else "B",
                wbs_affected=wbs_val,
                entry_date=coerce_diary_date(row[col_map.get("date", -1)]) if "date" in col_map else date.today(),
                line_section=str(row[col_map["section"]]).strip() if "section" in col_map and row[col_map["section"]] is not None else None,
                content=task_text,
                solution=str(row[col_map["solution"]]).strip() if "solution" in col_map and row[col_map["solution"]] is not None else None,
                physical_action=str(row[col_map["physical"]]).strip() if "physical" in col_map and row[col_map["physical"]] is not None else None,
                purpose=str(row[col_map["purpose"]]).strip() if "purpose" in col_map and row[col_map["purpose"]] is not None else None,
                author_name=str(row[col_map["author"]]).strip() if "author" in col_map and row[col_map["author"]] is not None else current_user.full_name,
                deadline_str=str(row[col_map["deadline"]]).strip() if "deadline" in col_map and row[col_map["deadline"]] is not None else None,
                status=str(row[col_map["status"]]).strip() if "status" in col_map and row[col_map["status"]] is not None else "В работе",
                source=str(row[col_map["source"]]).strip() if "source" in col_map and row[col_map["source"]] is not None else "Импорт Excel",
                entry_type=DiaryEntryType.DAILY_LOG,
            )
            db.add(entry)
            imported_count += 1

        await db.commit()
        return {
            "status": "success",
            "imported_count": imported_count,
            "message": f"Успешно импортировано {imported_count} записей в журнал производства работ",
        }
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)