"""
Unit-тесты для ExcelSmartTableParser (Legacy Migration).
"""

import os
from datetime import date
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
import pytest

from parser.src.excel_parser import ExcelSmartTableParser


@pytest.fixture
def sample_excel_file(tmp_path) -> str:
    """
    Генерирует временный Excel-файл с настоящими Умными таблицами (ListObject)
    и иероглифом '鐵' в имени файла для проверки кодировки.
    """
    file_path = str(tmp_path / "鐵__1.0.14h_test.xlsx")
    wb = openpyxl.Workbook()

    # --- Лист 1: Паспорт проекта ---
    ws_project = wb.active
    ws_project.title = "Project_Info"
    
    project_rows = [
        ["Параметр", "Значение"],
        ["Код проекта", "PRJ-STEEL-001"],
        ["Наименование", "Строительство Цеха Металлоконструкций (鐵)"],
        ["Бюджет", 150000000.00],
        ["Локация", "Индустриальный парк 'Север'"],
        ["Дата старта", "2026-09-01"],
        ["Дата финиша", "2026-12-31"],
    ]
    for row in project_rows:
        ws_project.append(row)

    # Создаем умную таблицу Tbl_Project_Data
    tab_proj = Table(displayName="Tbl_Project_Data", ref="A1:B7")
    tab_proj.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws_project.add_table(tab_proj)

    # --- Лист 2: Календарный план ---
    ws_tasks = wb.create_sheet(title="Schedule")
    task_rows = [
        ["WBS", "Наименование", "Длительность", "План старт", "План финиш", "Факт старт", "% выполнения", "Предшественники"],
        ["1.0", "Подготовительный этап", 5, "2026-09-01", "2026-09-05", "2026-09-01", 100.0, ""],
        ["1.1", "Земляные работы", 10, "2026-09-08", "2026-09-18", "2026-09-08", 50.0, "1.0FS+2"],
        ["2.0", "Монтаж металлокаркаса", 15, "2026-09-21", "2026-10-05", None, 0.0, "1.1SS+5, 1.0"],
        ["3.0", "Сдача объекта (Веха)", 0, "2026-10-06", "2026-10-06", None, 0.0, "2.0FF"],
    ]
    for row in task_rows:
        ws_tasks.append(row)

    # Создаем умную таблицу Tbl_Tasks_Schedule
    tab_tasks = Table(displayName="Tbl_Tasks_Schedule", ref="A1:H5")
    tab_tasks.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws_tasks.add_table(tab_tasks)

    wb.save(file_path)
    return file_path


def test_dependency_string_parsing():
    """Тест парсера сложных строк предшественников."""
    parser = ExcelSmartTableParser()

    deps1 = parser._parse_dependency_string("1.1FS+2, 2.3SS-1; 3.0FF")
    assert len(deps1) == 3
    assert deps1[0].predecessor_wbs == "1.1"
    assert deps1[0].dependency_type == "FS"
    assert deps1[0].lag_days == 2

    assert deps1[1].predecessor_wbs == "2.3"
    assert deps1[1].dependency_type == "SS"
    assert deps1[1].lag_days == -1

    assert deps1[2].predecessor_wbs == "3.0"
    assert deps1[2].dependency_type == "FF"
    assert deps1[2].lag_days == 0


def test_excel_smart_table_parser(sample_excel_file):
    """Тест полного цикла извлечения пакета проекта из файла Excel."""
    parser = ExcelSmartTableParser()
    package = parser.parse_file(sample_excel_file)

    # 1. Проверка паспорта
    assert package.passport.code == "PRJ-STEEL-001"
    assert "Строительство Цеха Металлоконструкций" in package.passport.name
    assert package.passport.planned_budget == 150000000.00
    assert package.passport.planned_start_date == date(2026, 9, 1)

    # 2. Проверка задач
    assert len(package.tasks) == 4

    t1 = package.tasks[0]
    assert t1.wbs_code == "1.0"
    assert t1.duration_days == 5
    assert t1.progress_percent == 100.0
    assert t1.actual_start == date(2026, 9, 1)
    assert len(t1.dependencies) == 0

    t2 = package.tasks[1]
    assert t2.wbs_code == "1.1"
    assert len(t2.dependencies) == 1
    assert t2.dependencies[0].predecessor_wbs == "1.0"
    assert t2.dependencies[0].lag_days == 2

    t3 = package.tasks[2]
    assert t3.wbs_code == "2.0"
    assert len(t3.dependencies) == 2
    assert t3.dependencies[0].predecessor_wbs == "1.1"
    assert t3.dependencies[0].dependency_type == "SS"
    assert t3.dependencies[0].lag_days == 5

    t4 = package.tasks[3]
    assert t4.wbs_code == "3.0"
    assert t4.duration_days == 0
    assert t4.dependencies[0].dependency_type == "FF"