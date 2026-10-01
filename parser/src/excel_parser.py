"""
Парсер умных таблиц Excel (Tbl_Project_*, Tbl_Tasks_*) для миграции legacy-данных.
"""

import os
import re
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Tuple
import openpyxl


class ExcelParsingError(Exception):
    """Ошибка при чтении или парсинге Excel файла."""
    pass


class ExcelSmartTableParser:
    """
    Извлекает данные паспорта проекта и задач из умных таблиц Excel.
    """

    DEP_PATTERN = re.compile(
        r"^(?P<wbs>[\d\.]+)(?P<type>FS|SS|FF|SF)?(?P<sign>[\+\-])?(?P<lag>\d+)?$",
        re.IGNORECASE,
    )

    def parse_file(self, file_path: str):
        """
        Главная точка входа: открывает Excel-файл и парсит паспорт и задачи.
        """
        from parser.src.schema import ParsedProjectPackage

        if not os.path.exists(file_path):
            raise ExcelParsingError(f"Файл не найден: {file_path}")

        filename = os.path.basename(file_path)

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, keep_vba=True)
        except Exception as e:
            raise ExcelParsingError(f"Не удалось открыть Excel-файл {filename}: {str(e)}")

        try:
            # 1. Ищем таблицу задач по всем реальным именам из базы:
            # tblM06Data, _03_M06_Данные, График_проекта, _02_ПортфельЭтапов, Tbl_Tasks, Tbl_Schedule
            tasks_table = None
            tasks_ws = None

            KNOWN_TASK_PREFIXES = (
                "tblm06data", "_03_m06", "график_проекта", "_02_портфельэтапов",
                "tbl_tasks", "tbl_schedule", "tblm07view", "этапы", "гант"
            )

            for prefix in KNOWN_TASK_PREFIXES:
                t, ws = self._find_smart_table(wb, prefix=prefix)
                if t:
                    tasks_table, tasks_ws = t, ws
                    break

            # Если явный ListObject не найден, ищем лист '03_M06_Данные' или 'План-график' напрямую
            if not tasks_table:
                for sheetname in wb.sheetnames:
                    lower_sheet = sheetname.lower()
                    if any(k in lower_sheet for k in ("03_m06", "план-график", "данные", "портфельэтапов")):
                        tasks_ws = wb[sheetname]
                        # Создаём псевдо-диапазон по всем заполненным строкам листа
                        tasks_table = f"A1:{openpyxl.utils.get_column_letter(tasks_ws.max_column)}{tasks_ws.max_row}"
                        break

            if not tasks_table or not tasks_ws:
                raise ExcelParsingError(
                    f"В файле {filename} не найден лист или таблица задач (ожидались: '03_M06_Данные', 'План-график' или 'График_проекта')"
                )

            # 2. Ищем паспорт проекта (Проекты, MinimalPassportTable, PassportTable18, Tbl_Project)
            project_table = None
            project_ws = None
            for prefix in ("minimalpassporttable", "passporttable", "tblproject", "проекты", "tbl_project"):
                p_t, p_ws = self._find_smart_table(wb, prefix=prefix)
                if p_t:
                    project_table, project_ws = p_t, p_ws
                    break

            passport = self._parse_passport(project_table, project_ws, fallback_code=filename)
            tasks = self._parse_tasks(tasks_table, tasks_ws)

            return ParsedProjectPackage(
                passport=passport,
                tasks=tasks,
                source_filename=filename,
            )
        finally:
            wb.close()

    def _find_smart_table(
        self, wb: openpyxl.Workbook, prefix: str
    ) -> Tuple[Any, Optional[openpyxl.worksheet.worksheet.Worksheet]]:
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            tables = getattr(ws, "tables", {})
            if hasattr(tables, "items"):
                for table_name, table in tables.items():
                    if str(table_name).lower().startswith(prefix.lower()):
                        return table, ws
            elif isinstance(tables, list):
                for table in tables:
                    name = getattr(table, "name", getattr(table, "displayName", ""))
                    if str(name).lower().startswith(prefix.lower()):
                        return table, ws
        return None, None

    def _get_table_rows(
        self, table: Any, ws: openpyxl.worksheet.worksheet.Worksheet
    ) -> List[Dict[str, Any]]:
        ref = getattr(table, "ref", table if isinstance(table, str) else None)
        if not ref:
            return []

        data_range = ws[ref]
        
        if not isinstance(data_range, tuple):
            data_range = ((data_range,),)
        elif data_range and not isinstance(data_range[0], tuple):
            data_range = (data_range,)

        if not data_range:
            return []

        headers = [
            str(cell.value).strip() if cell.value is not None else f"col_{idx}" 
            for idx, cell in enumerate(data_range[0])
        ]
        
        rows = []
        for row_cells in data_range[1:]:
            row_dict = {}
            for header, cell in zip(headers, row_cells):
                row_dict[header] = cell.value
            if any(v is not None for v in row_dict.values()):
                rows.append(row_dict)
        return rows

    def _get_first_value(self, d: Dict[str, Any], keys: List[str], default: Any = None) -> Any:
        """Безопасное извлечение первого не-None значения с сохранением нулей (0)."""
        for k in keys:
            if k in d and d[k] is not None and str(d[k]).strip() != "":
                return d[k]
        return default

    def _parse_passport(
        self,
        table: Any,
        ws: Optional[openpyxl.worksheet.worksheet.Worksheet],
        fallback_code: str,
    ):
        from parser.src.schema import ParsedProjectPassport

        default_start = date.today()
        default_finish = date.today()

        if not table or not ws:
            clean_name = os.path.splitext(fallback_code)[0]
            return ParsedProjectPassport(
                code=clean_name[:64],
                name=clean_name,
                planned_start_date=default_start,
                planned_finish_date=default_finish,
            )

        rows = self._get_table_rows(table, ws)
        if not rows:
            clean_name = os.path.splitext(fallback_code)[0]
            return ParsedProjectPassport(
                code=clean_name[:64],
                name=clean_name,
                planned_start_date=default_start,
                planned_finish_date=default_finish,
            )

        def safe_json_val(v: Any) -> Any:
            if v is None:
                return None
            if isinstance(v, (datetime, date)):
                return v.isoformat()
            if isinstance(v, (int, float, bool, str)):
                return v
            return str(v)

        flat_data: Dict[str, Any] = {}
        for r in rows:
            values = list(r.values())
            if len(values) >= 2 and values[0] is not None:
                flat_data[str(values[0]).strip().lower()] = safe_json_val(values[1])
            else:
                for k, v in r.items():
                    flat_data[k.strip().lower()] = safe_json_val(v)

        code = str(self._get_first_value(flat_data, ["code", "код проекта", "код"], default=fallback_code))[:64]
        name = str(self._get_first_value(flat_data, ["name", "наименование", "название проекта"], default=code))
        
        start_date = self._coerce_date(
            self._get_first_value(flat_data, ["planned_start", "дата старта", "старт", "начало"]),
            default=default_start,
        )
        finish_date = self._coerce_date(
            self._get_first_value(flat_data, ["planned_finish", "дата финиша", "финиш", "окончание"]),
            default=default_finish,
        )

        return ParsedProjectPassport(
            code=code,
            name=name,
            description=self._get_first_value(flat_data, ["description", "описание"]),
            location=self._get_first_value(flat_data, ["location", "локация"]),
            currency=str(self._get_first_value(flat_data, ["currency", "валюта"], default="RUB")),
            planned_budget=self._coerce_float(self._get_first_value(flat_data, ["budget", "бюджет"])),
            planned_start_date=start_date,
            planned_finish_date=finish_date,
            passport_data=flat_data,
        )

    def _parse_tasks(
        self, table: Any, ws: openpyxl.worksheet.worksheet.Worksheet
    ):
        from parser.src.schema import ParsedTask

        def safe_json_val(v: Any) -> Any:
            if v is None:
                return None
            if isinstance(v, (datetime, date)):
                return v.isoformat()
            if isinstance(v, (int, float, bool, str)):
                return v
            return str(v)

        rows = self._get_table_rows(table, ws)
        tasks: List[ParsedTask] = []

        for row in rows:
            norm_row = {str(k).strip().lower(): safe_json_val(v) for k, v in row.items()}

            wbs = str(self._get_first_value(
                norm_row, ["wbs", "код сдр", "код wbs", "№ п/п", "id этапа", "id", "порядок"], default=""
            )).strip()

            if not wbs:
                continue

            name = str(self._get_first_value(
                norm_row, ["этап/задача", "этап", "наименование", "name", "наименование работ", "задача"], default=f"Этап {wbs}"
            )).strip()

            raw_duration = self._get_first_value(
                norm_row, ["duration", "длительность", "длительность (дней)"], default=1
            )
            duration = self._coerce_int(raw_duration, default=1)

            p_start = self._coerce_date(
                self._get_first_value(norm_row, ["planned_start", "план старт", "начало"]),
                default=date.today(),
            )
            p_finish = self._coerce_date(
                self._get_first_value(norm_row, ["planned_finish", "план финиш", "окончание"]),
                default=p_start,
            )

            a_start = self._coerce_date(
                self._get_first_value(norm_row, ["actual_start", "факт старт"])
            )
            a_finish = self._coerce_date(
                self._get_first_value(norm_row, ["actual_finish", "факт финиш"])
            )
            progress = self._coerce_float(
                self._get_first_value(norm_row, ["progress", "% выполнения", "прогресс"]),
                default=0.0,
            )

            raw_deps = str(self._get_first_value(
                norm_row, ["predecessors", "предшественники", "связи"], default=""
            ))
            dependencies = self._parse_dependency_string(raw_deps)

            # Автоматический захват Pyrus-ссылок из любых столбцов (K, L, S, pyrus, url, гиперссылок)
            pyrus_url = None
            for col_key, col_val in norm_row.items():
                if any(x in col_key for x in ("pyrus", "пайрус", "ссылк", "url", "task_url")):
                    if col_val and "pyrus" in str(col_val).lower():
                        pyrus_url = str(col_val).strip()
                        break
            if not pyrus_url:
                for col_key, col_val in norm_row.items():
                    if col_val and isinstance(col_val, str) and "pyrus.com" in col_val:
                        m = re.search(r"https?://[^\s\"'<>]*pyrus[^\s\"'<>]*", col_val)
                        if m:
                            pyrus_url = m.group(0)
                            break

            custom_attrs = dict(norm_row)
            if pyrus_url:
                custom_attrs["pyrus_url"] = pyrus_url

            tasks.append(
                ParsedTask(
                    wbs_code=wbs,
                    name=name,
                    duration_days=duration,
                    planned_start=p_start,
                    planned_finish=p_finish,
                    actual_start=a_start,
                    actual_finish=a_finish,
                    progress_percent=progress,
                    raci_assignee=self._get_first_value(norm_row, ["responsible", "ответственный"]),
                    dependencies=dependencies,
                    custom_attributes=custom_attrs,
                )
            )

        return tasks

    def _parse_dependency_string(self, raw_str: str):
        from parser.src.schema import ParsedDependency

        if not raw_str or raw_str.strip() in ("", "None", "nan"):
            return []

        deps: List[ParsedDependency] = []
        parts = [p.strip() for p in re.split(r"[,;]\s*", raw_str) if p.strip()]

        for part in parts:
            match = self.DEP_PATTERN.match(part)
            if match:
                pred_wbs = match.group("wbs")
                dep_type = (match.group("type") or "FS").upper()
                sign = match.group("sign") or "+"
                lag_val = int(match.group("lag") or 0)
                lag_days = lag_val if sign == "+" else -lag_val

                deps.append(
                    ParsedDependency(
                        predecessor_wbs=pred_wbs,
                        dependency_type=dep_type,
                        lag_days=lag_days,
                    )
                )

        return deps

    @staticmethod
    def _coerce_date(val: Any, default: Optional[date] = None) -> Optional[date]:
        if val is None or str(val).strip() in ("", "None", "nan"):
            return default
        if isinstance(val, datetime):
            return val.date()
        if isinstance(val, date):
            return val
        if isinstance(val, str):
            for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%Y"):
                try:
                    return datetime.strptime(val.strip(), fmt).date()
                except ValueError:
                    pass
        return default

    @staticmethod
    def _coerce_int(val: Any, default: int = 0) -> int:
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return default

    @staticmethod
    def _coerce_float(val: Any, default: Optional[float] = None) -> Optional[float]:
        try:
            return float(val)
        except (ValueError, TypeError):
            return default